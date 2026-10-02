"""Render Slide specs to a python-pptx Presentation (16:9, 13.333 x 7.5 in)."""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.oxml import parse_xml

from .design import (W, H, Box, Line, Ellipse, TextSpec)


def _em(v):
    return Inches(round(v, 4))


def _hex(c):
    return RGBColor.from_string(c)


_ALIGN = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}
_ANCHOR = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}


def _apply_gradient(fill, top, bottom, ang=0):
    fill.gradient()
    stops = fill.gradient_stops
    stops[0].color.rgb = _hex(top)
    stops[1].color.rgb = _hex(bottom)
    fill.gradient_angle = ang


def _set_bg(slide, spec):
    if spec.bg_grad:
        top, bottom, ang = spec.bg_grad
        _apply_gradient(slide.background.fill, top, bottom, ang)
    elif spec.bg:
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = _hex(spec.bg)


def _add_box(slide, b):
    st = MSO_SHAPE.ROUNDED_RECTANGLE if b.round_ else MSO_SHAPE.RECTANGLE
    sp = slide.shapes.add_shape(st, _em(b.x), _em(b.y), _em(b.w), _em(b.h))
    if b.gradient:
        _apply_gradient(sp.fill, b.gradient[0], b.gradient[1],
                        b.gradient[2] if len(b.gradient) > 2 else 0)
    elif b.fill:
        sp.fill.solid()
        sp.fill.fore_color.rgb = _hex(b.fill)
    else:
        sp.fill.background()
    if b.line:
        sp.line.color.rgb = _hex(b.line)
        sp.line.width = Pt(b.line_w)
    else:
        sp.line.fill.background()
    sp.shadow.inherit = False
    return sp


def _add_line(slide, ln):
    (x1, y1), (x2, y2) = ln.endpoints
    conn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,
                                      _em(x1), _em(y1), _em(x2), _em(y2))
    conn.line.color.rgb = _hex(ln.color)
    conn.line.width = Pt(ln.width)
    conn.shadow.inherit = False
    return conn


def _add_ellipse(slide, e):
    sp = slide.shapes.add_shape(MSO_SHAPE.OVAL, _em(e.x), _em(e.y), _em(e.w), _em(e.h))
    if e.fill:
        sp.fill.solid()
        sp.fill.fore_color.rgb = _hex(e.fill)
    else:
        sp.fill.background()
    if e.line:
        sp.line.color.rgb = _hex(e.line)
        sp.line.width = Pt(e.line_w)
    else:
        sp.line.fill.background()
    sp.shadow.inherit = False
    return sp


def _apply_run(prun, drun):
    prun.font.size = Pt(drun.size)
    prun.font.bold = drun.bold
    prun.font.italic = drun.italic
    if drun.font:
        prun.font.name = drun.font
    prun.font.color.rgb = _hex(drun.color)


def _set_bullet(p, color):
    """Native em-dash bullet (buChar) matching the delivered deck."""
    A = "http://schemas.openxmlformats.org/drawingml/2006/main"
    pPr = p._p.get_or_add_pPr()
    pPr.set("marL", "201600")
    pPr.set("indent", "-201600")
    for xml in ('<a:buClr xmlns:a="%s"><a:srgbClr val="%s"/></a:buClr>' % (A, color),
                '<a:buSzPct xmlns:a="%s" val="76000"/>' % A,
                '<a:buFont xmlns:a="%s" typeface="Arial"/>' % A,
                '<a:buChar xmlns:a="%s" char="\u2014"/>' % A):
        pPr.append(parse_xml(xml))


def _add_text(slide, ts):
    tb = slide.shapes.add_textbox(_em(ts.x), _em(ts.y), _em(ts.w), _em(ts.h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = _ANCHOR.get(ts.anchor, MSO_ANCHOR.TOP)
    first = True
    for para in ts.paragraphs:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = _ALIGN.get(para.align, PP_ALIGN.LEFT)
        p.space_before = Pt(para.before)
        p.space_after = Pt(para.after)
        p.line_spacing = para.leading
        if para.bullet:
            _set_bullet(p, para.bullet_color or "A67F3B")
        for drun in para.runs:
            prun = p.add_run()
            prun.text = drun.text
            _apply_run(prun, drun)
    return tb


def _add_element(slide, el):
    if isinstance(el, Box):
        return _add_box(slide, el)
    if isinstance(el, Line):
        return _add_line(slide, el)
    if isinstance(el, Ellipse):
        return _add_ellipse(slide, el)
    if isinstance(el, TextSpec):
        return _add_text(slide, el)
    raise TypeError("unknown element: %r" % type(el))


def render_slide(prs, spec):
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank layout
    _set_bg(slide, spec)
    for el in spec.elements:
        _add_element(slide, el)
    return slide


def render(prs, specs):
    prs.slide_width = Inches(W)
    prs.slide_height = Inches(H)
    for spec in specs:
        render_slide(prs, spec)
    return prs


def new_presentation():
    return Presentation()
