"""Compare a rebuilt PPTX with a reference, including editable visual styles.

Usage:
    python -m build.engine.verify reference.pptx rebuilt.pptx [--tol 0.03]
"""
import argparse
from collections import Counter
from pptx import Presentation
from pptx.util import Emu

A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


def _inch(value):
    return value / 914400.0


def _color(node):
    if node is None:
        return None
    rgb = node.find(".//" + A + "srgbClr")
    if rgb is not None:
        return "rgb:" + rgb.get("val", "")
    scheme = node.find(".//" + A + "schemeClr")
    if scheme is not None:
        return "scheme:" + scheme.get("val", "")
    return None


def _fill(shape):
    sp = shape._element.find(".//" + A + "spPr")
    if sp is None:
        return None
    solid = sp.find("./" + A + "solidFill")
    if solid is not None:
        return ("solid", _color(solid))
    grad = sp.find("./" + A + "gradFill")
    if grad is not None:
        stops = tuple((g.get("pos"), _color(g)) for g in grad.findall(".//" + A + "gs"))
        lin = grad.find(".//" + A + "lin")
        angle = lin.get("ang") if lin is not None else None
        return ("gradient", stops, angle)
    return None


def _line(shape):
    sp = shape._element.find(".//" + A + "spPr")
    ln = sp.find("./" + A + "ln") if sp is not None else None
    if ln is None:
        return None
    width = ln.get("w")
    dash = ln.find("./" + A + "prstDash")
    return (_color(ln.find("./" + A + "solidFill")),
            int(width) / 12700.0 if width else None,
            ln.find("./" + A + "noFill") is not None,
            dash.get("val") if dash is not None else None)


def _text_styles(shape):
    if not shape.has_text_frame:
        return None
    result = []
    for para in shape.text_frame.paragraphs:
        pp = para._p.find(A + "pPr")

        def attr(name):
            return pp.get(name) if pp is not None else None

        def val(path):
            node = pp.find(path) if pp is not None else None
            return node.get("val") if node is not None else None

        bullet = pp.find(A + "buChar") if pp is not None else None
        runs = []
        for run in para.runs:
            font = run.font
            try:
                color = str(font.color.rgb) if font.color and font.color.rgb else None
            except Exception:
                color = None
            runs.append((run.text, font.name,
                         float(font.size.pt) if font.size else None,
                         font.bold, font.italic, color))
        result.append((attr("algn"), attr("marL"), attr("indent"),
                       bullet.get("char") if bullet is not None else None,
                       val(".//" + A + "spcPct"),
                       val("./" + A + "spcBef/" + A + "spcPts"),
                       val("./" + A + "spcAft/" + A + "spcPts"),
                       tuple(runs)))
    return tuple(result)


def _record(shape):
    geom = tuple(round(_inch(v), 3) for v in
                 (shape.left, shape.top, shape.width, shape.height))
    text = " ".join(shape.text_frame.text.split()) if shape.has_text_frame else ""
    return {"type": str(shape.shape_type), "geom": geom, "text": text,
            "fill": _fill(shape), "line": _line(shape),
            "text_styles": _text_styles(shape)}


def _distance(a, b):
    return max(abs(x - y) for x, y in zip(a["geom"], b["geom"]))


def compare(reference, candidate, tolerance=0.03, verbose=True):
    ref = Presentation(reference)
    got = Presentation(candidate)
    counts = Counter()
    examples = []
    if len(ref.slides) != len(got.slides):
        counts["slide_count"] = abs(len(ref.slides) - len(got.slides))
    for si in range(min(len(ref.slides), len(got.slides))):
        left = [_record(s) for s in ref.slides[si].shapes]
        right = [_record(s) for s in got.slides[si].shapes]
        used = set()
        for source in left:
            candidates = []
            for j, rebuilt in enumerate(right):
                if j in used or source["type"] != rebuilt["type"]:
                    continue
                distance = _distance(source, rebuilt)
                if distance <= tolerance:
                    # Prefer exact text when choosing among overlapping shapes.
                    text_penalty = 0 if source["text"] == rebuilt["text"] else 1
                    candidates.append((text_penalty, distance, j, rebuilt))
            if not candidates:
                counts["missing"] += 1
                if len(examples) < 10:
                    examples.append((si + 1, "missing", source))
                continue
            _, _, j, rebuilt = min(candidates)
            used.add(j)
            for key in ("text", "fill", "line", "text_styles"):
                if source[key] != rebuilt[key]:
                    counts[key] += 1
                    if len(examples) < 10:
                        examples.append((si + 1, key, source, rebuilt))
        counts["extra"] += sum(1 for j in range(len(right)) if j not in used)
    counts += Counter()  # discard zero-value keys accumulated by Counter arithmetic
    if verbose:
        print("slides: reference=%d rebuilt=%d" % (len(ref.slides), len(got.slides)))
        print("geometry tolerance: %.3f in" % tolerance)
        print("differences:", dict(counts) or "none")
        for example in examples:
            si, kind = example[:2]
            print("  slide %03d %s" % (si, kind))
    return counts


def main(argv=None):
    ap = argparse.ArgumentParser(description="Compare a rebuilt deck to its reference.")
    ap.add_argument("reference")
    ap.add_argument("candidate")
    ap.add_argument("--tol", type=float, default=0.03,
                    help="geometry tolerance in inches (default: 0.03)")
    args = ap.parse_args(argv)
    counts = compare(args.reference, args.candidate, args.tol)
    return 1 if counts else 0


if __name__ == "__main__":
    raise SystemExit(main())
